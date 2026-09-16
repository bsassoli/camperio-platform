# -*- coding: utf-8 -*-
"""Report Variazioni: l'euro e' il dato principale, la valuta del titolo la colonna a fianco.

Decisione dell'autore del 16/09/2026 (voce 8 di CENSIMENTO-NUOVA-VERSIONE.md): il dato
principale resta in euro con il cambio incluso, la variazione nella valuta del titolo si
aggiunge accanto, cosi' l'effetto cambio si legge per differenza. Medie, conteggi e
ordinamento continuano a seguire l'euro.
"""
import openpyxl
import pytest
from docx import Document

import data_layer as DL
import lookthrough as L


_PF = {"meta": {"schema": "ANTASIMGEST", "codcli": "DEMO01",
                "data_prec": "2026-09-01", "data": "2026-09-08", "nav": 1_000_000.0}}

# Titolo in USD: prezzo locale +10%, ma il dollaro si indebolisce (1,00 -> 1,20 EUR/USD),
# quindi in euro il titolo perde. E' il caso che rende visibile l'effetto cambio.
# Titolo in EUR: nessun cambio di mezzo, le due colonne devono coincidere.
_ROWS = [
    {"nome": "USD SU FX GIU", "ccy": "USD", "loc_dal": 100.0, "loc_al": 110.0,
     "fx_dal": 1.0, "fx_al": 1.2, "qty_al": 1000.0},
    {"nome": "EURO PIATTO", "ccy": "EUR", "loc_dal": 50.0, "loc_al": 55.0,
     "fx_dal": 1.0, "fx_al": 1.0, "qty_al": 1000.0},
    {"nome": "SENZA PREZZO", "ccy": "USD", "loc_dal": None, "loc_al": None,
     "fx_dal": 1.0, "fx_al": 1.2, "qty_al": 1000.0},
]

_VAR_USD_EUR = 110.0 / 1.2 / 100.0 - 1     # -8,33%: in euro scende
_VAR_USD_LOC = 110.0 / 100.0 - 1           # +10,0%: in valuta sale


@pytest.fixture()
def d(tmp_path, monkeypatch):
    monkeypatch.setattr(DL, "price_changes",
                        lambda *a, **k: {"nav": 1_000_000.0, "dal": "2026-09-01",
                                         "al": "2026-09-08", "rows": _ROWS})
    return L.build_variazioni(_PF, str(tmp_path))


def _riga(d, prefisso):
    return next(r for r in d["rows"] if r["name"].lower().startswith(prefisso))


def test_le_due_variazioni_convivono_e_il_cambio_si_legge_per_differenza(d):
    usd = _riga(d, "usd su fx")
    assert usd["var"] == pytest.approx(_VAR_USD_EUR)
    assert usd["var_loc"] == pytest.approx(_VAR_USD_LOC)
    assert usd["var"] < 0 < usd["var_loc"]          # senza la colonna l'effetto cambio sparirebbe


def test_titolo_in_euro_le_due_colonne_coincidono(d):
    eur = _riga(d, "euro piatto")
    assert eur["var"] == pytest.approx(0.10)
    assert eur["var_loc"] == pytest.approx(eur["var"])


def test_medie_conteggi_e_ordinamento_seguono_l_euro(d):
    vars_eur = [r["var"] for r in d["rows"] if r["var"] is not None]
    assert vars_eur == sorted(vars_eur, reverse=True)
    assert d["n"] == 2 and d["up"] == 1 and d["down"] == 1
    assert d["media"] == pytest.approx((0.10 + _VAR_USD_EUR) / 2)
    assert d["best"]["var"] == pytest.approx(0.10)                  # il titolo in euro
    assert d["worst"]["var"] == pytest.approx(_VAR_USD_EUR)         # il titolo in dollari


def test_riga_senza_prezzo_non_inventa_variazioni(d):
    r = _riga(d, "senza prezzo")
    assert r["var"] is None and r["var_loc"] is None


def test_html_espone_entrambe_le_colonne(d):
    html = L.variazioni_html(d)
    assert "Var. % €" in html and "Var. % valuta" in html
    assert "valuta del titolo" in html          # il sottotitolo spiega come leggere la differenza
    assert "n.d." in html                       # la riga senza prezzo resta dichiarata


def test_excel_espone_entrambe_le_colonne(d, tmp_path):
    path = str(tmp_path / "variazioni.xlsx")
    L.variazioni_excel(d, path)
    ws = openpyxl.load_workbook(path).active
    intestazioni = next(([c.value for c in row] for row in ws.iter_rows() if row[0].value == "#"), None)
    assert intestazioni is not None, "riga di intestazione non trovata"
    assert "Var. % settimana (€)" in intestazioni
    assert "Var. % settimana (valuta)" in intestazioni
    col_loc = intestazioni.index("Var. % settimana (valuta)") + 1
    valori = [ws.cell(row=r, column=col_loc).value for r in range(1, ws.max_row + 1)]
    assert any(v == pytest.approx(_VAR_USD_LOC) for v in valori if isinstance(v, float))


def test_word_espone_entrambe_le_colonne(d, tmp_path):
    path = str(tmp_path / "variazioni.docx")
    L.variazioni_word(d, path)
    t = Document(path).tables[0]
    intestazioni = [c.text for c in t.rows[0].cells]
    assert "Var. % €" in intestazioni and "Var. % val." in intestazioni
